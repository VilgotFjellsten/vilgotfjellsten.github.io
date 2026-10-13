"""Export podcast rows to CSV and/or Excel (.xlsx)."""

from __future__ import annotations

import csv
import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from podcast_scraper.models import PodcastRow

logger = logging.getLogger(__name__)

COLUMNS = [
    "Podcast name",
    "Niche / sub-niche",
    "Description",
    "Email (official)",
    "Email (host-derived)",
    "Contact form URL",
    "Website",
    "Socials",
    "Host name(s)",
    "Platforms found on",
    "Source URLs",
    "RSS URL",
    "Apple ID",
    "Podcast Index ID",
    "Confidence / notes",
    "Do not contact",
]


def export_csv(rows: list[PodcastRow], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=COLUMNS, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row.to_export_dict())
    logger.info("Wrote CSV %s (%d rows)", path, len(rows))
    return path


def export_xlsx(rows: list[PodcastRow], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Podcasts"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        data = row.to_export_dict()
        ws.append([data.get(c, "") for c in COLUMNS])
    # Reasonable column widths
    for idx, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = min(
            40, max(12, len(col) + 2)
        )
    wb.save(path)
    logger.info("Wrote Excel %s (%d rows)", path, len(rows))
    return path


def export(
    rows: list[PodcastRow],
    output_dir: Path,
    *,
    job_id: str,
    formats: tuple[str, ...] = ("csv", "xlsx"),
) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    stem = f"{job_id}_podcasts"
    if "csv" in formats:
        written.append(export_csv(rows, output_dir / f"{stem}.csv"))
    if "xlsx" in formats:
        written.append(export_xlsx(rows, output_dir / f"{stem}.xlsx"))
    return written
